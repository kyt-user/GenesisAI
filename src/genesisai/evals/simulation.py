"""P7 隔离产品仿真 Runner 与机器可读报告。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import shutil
import stat
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import yaml

from genesisai.memory.manager import MemoryManager
from genesisai.model.config import build_model
from genesisai.agent.runner import Runner
from genesisai.shared.security import Access
from genesisai.skills.registry import SkillRegistry
from genesisai.state.store import Store, atomic_json
from genesisai.runtime.tool_runtime import ToolRuntime


CASE_FIELDS = frozenset({"id", "name", "stages", "prompt", "fixture", "assertions", "forbidden", "artifacts"})


def inventory(root):
    return {
        path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(Path(root).rglob("*")) if path.is_file()
    }


def load_dataset(root):
    root = Path(root).resolve()
    manifest = yaml.safe_load((root / "manifest.yaml").read_text(encoding="utf-8"))
    if not isinstance(manifest, dict) or set(manifest) != {"cases"} or not isinstance(manifest["cases"], list):
        raise ValueError("P7 manifest.yaml 必须且只能包含 cases 数组")
    cases = []
    for identifier in manifest["cases"]:
        case_root = root / identifier
        value = yaml.safe_load((case_root / "case.yaml").read_text(encoding="utf-8"))
        if not isinstance(value, dict) or set(value) != CASE_FIELDS or value["id"] != identifier:
            raise ValueError(f"P7 case 清单字段无效：{identifier}")
        fixture = (case_root / value["fixture"]).resolve(strict=True)
        if not fixture.is_relative_to(case_root) or not fixture.is_dir():
            raise ValueError(f"P7 fixture 越界：{identifier}")
        cases.append((value, fixture))
    return cases


class FixtureProvider:
    name = "fixture"
    def __init__(self, data): self.data = data
    def search(self, query, limit=5): return self.data["results"][:limit]


class FixtureNetwork:
    def __init__(self, data): self.data = data
    def fetch(self, url):
        if url != self.data["page"]["url"]: raise ValueError("fixture URL 不存在")
        return {**self.data["page"], "truncated": False}


def make_runtime(workspace, run_root, *, providers=None, network=None):
    state = run_root / "state"; output = run_root / "outputs"; output.mkdir(parents=True, exist_ok=True)
    store = Store(state, workspace=workspace)
    store.data.update(grants=[str(workspace)], output=str(output), run_id="run_simulation")
    store.save()
    runtime = ToolRuntime(store, Access([workspace], output, state, workspace_root=workspace), providers=providers, network=network, confirm_writes=False, confirm_search=False, confirm_shell=False)
    return store, runtime, output


def call(runtime, name, args, suffix=""):
    ident = f"sim_{name}_{len(runtime.store.data['calls'])}{suffix}"
    return runtime.execute({"id": ident, "name": name, "arguments": json.dumps(args, ensure_ascii=False)})


def require_ok(result):
    if not result.get("ok"):
        raise RuntimeError((result.get("error") or {}).get("message", "工具失败"))
    return result["data"]


def runtime_report(store, runtime):
    """返回不含 Prompt 正文、工具参数和密钥的运行摘要。"""
    data = store.data
    run = data.get("run_runtime", {})
    return {
        "session_id": store.id,
        "run_id": data.get("run_id"),
        "runner_state": data.get("status"),
        "lifecycle": run.get("lifecycle"),
        "stop_reason": run.get("stop_reason"),
        "rounds": data.get("rounds", 0),
        "model_calls": (run.get("usage") or {}).get("model_calls", 0),
        "tool_calls": len(data.get("calls", {})),
        "tokens": (data.get("usage") or {}).get("total_tokens", 0),
        "prompt_hashes": run.get("prompt_hashes", {}),
        "active_tools": list(data.get("tool_runtime", {}).get("active_tools", [])),
        "active_skills": list(data.get("tool_runtime", {}).get("active_skills", [])),
        "permissions": {
            "network": "ask" if runtime.confirm_search else "allow",
            "writes": "ask" if runtime.confirm_writes else "allow",
            "shell": "ask" if runtime.confirm_shell else "allow",
        },
        "recovery_counts": run.get("recovery_counts", {}),
        "calls": [
            {"id": identifier, "name": item.get("call", {}).get("name"), "state": item.get("state")}
            for identifier, item in data.get("calls", {}).items()
        ],
        "changes": data.get("changes", []),
        "artifacts": list(data.get("artifacts", {})),
        "sources": list(data.get("sources", {})),
    }


def run_offline_case(case, workspace, run_root):
    identifier = case["id"]
    store, runtime, output = make_runtime(workspace, run_root)
    checks = {}
    if identifier == "case_01_code_fix":
        subprocess.run(["git", "init"], cwd=workspace, check=True, capture_output=True)
        subprocess.run(["git", "add", "."], cwd=workspace, check=True, capture_output=True)
        subprocess.run(["git", "-c", "user.name=GenesisAI", "-c", "user.email=simulation@example.invalid", "commit", "-m", "fixture"], cwd=workspace, check=True, capture_output=True)
        runtime.load_tools(["file_read", "file_patch", "test_run", "git_status", "git_diff"])
        original_test = (workspace / "test_calculator.py").read_bytes()
        before = require_ok(call(runtime, "file_read", {"path": "calculator.py"}))["sha256"]
        patch = "@@ -1,2 +1,2 @@\n def add(left, right):\n-    return left - right\n+    return left + right"
        require_ok(call(runtime, "file_patch", {"path": "calculator.py", "patch": patch, "expected_sha256": before}))
        test = require_ok(call(runtime, "test_run", {"command": [sys.executable, "-m", "pytest", "-q"], "cwd": ".", "timeout_seconds": 60}))
        diff = require_ok(call(runtime, "git_diff", {"cwd": "."}))
        checks = {"product_fixed": "left + right" in (workspace / "calculator.py").read_text(), "test_passed": test["exit_code"] == 0, "test_unchanged": (workspace / "test_calculator.py").read_bytes() == original_test, "git_read": "calculator.py" in diff["text"]}
    elif identifier == "case_02_file_management":
        (workspace / "windows.txt").write_bytes(b"\xef\xbb\xbfalpha\r\nbeta\r\n")
        runtime.load_tools(["file_read", "file_patch", "directory_create", "file_move", "file_create"])
        read = require_ok(call(runtime, "file_read", {"path": "windows.txt", "line_start": 1, "line_end": 2}))
        require_ok(call(runtime, "file_patch", {"path": "windows.txt", "patch": "@@ -1,2 +1,2 @@\n-alpha\n+ALPHA\n beta", "expected_sha256": read["sha256"]}))
        require_ok(call(runtime, "directory_create", {"path": "organized"}))
        require_ok(call(runtime, "file_move", {"source": "notes.md", "destination": "organized/notes.md"}))
        require_ok(call(runtime, "file_create", {"path": "summary.json", "content": "{\"organized\": true}", "source_refs": [], "scope": "workspace"}))
        binary = call(runtime, "file_patch", {"path": "binary.bin", "patch": "@@ -1,1 +1,1 @@\n-old\n+new"}, "binary")
        payload = (workspace / "windows.txt").read_bytes()
        checks = {"bom": payload.startswith(b"\xef\xbb\xbf"), "crlf": b"\r\n" in payload and b"ALPHA" in payload, "moved": (workspace / "organized" / "notes.md").is_file(), "created": (workspace / "summary.json").is_file(), "binary_rejected": not binary["ok"]}
    elif identifier == "case_03_web_research":
        data = json.loads((workspace / "public_pages.json").read_text(encoding="utf-8"))
        store, runtime, output = make_runtime(workspace, run_root / "web", providers=[FixtureProvider(data)], network=FixtureNetwork(data))
        runtime.load_tools(["search_query", "search_fetch"])
        query = require_ok(call(runtime, "search_query", {"query": "Apple event today"}))
        fetched = require_ok(call(runtime, "search_fetch", {"url": query["hits"][0]["url"]}))
        state = store.data["run_runtime"]
        checks = {"candidate": state["candidate_count"] == 1, "fetched": state["fetched_count"] == 1, "evidence": bool(state["evidence_refs"]), "source": bool(fetched["source_refs"]), "session_allow": not runtime.confirm_search}
    elif identifier == "case_04_memory_and_skill":
        runtime.load_tools(["test_run"])
        skill = SkillRegistry(workspace, runtime.registry).describe("project_check", include_content=True)
        test = require_ok(call(runtime, "test_run", {"command": [sys.executable, "-m", "pytest", "-q"], "cwd": ".", "timeout_seconds": 60}))
        first = MemoryManager(workspace)
        remembered = first.search("pytest", type="command")
        fact = first.add(type="project", title="项目配置", summary="项目测试入口", content="测试入口来自 project.toml", source_path="project.toml", verified=True)
        restarted = MemoryManager(workspace)
        (workspace / "project.toml").write_text('test_command = "changed"\n', encoding="utf-8")
        invalidated = restarted.validate_sources()
        checks = {"project_skill": skill["source"] == "project", "test_passed": test["exit_code"] == 0, "command_memory": bool(remembered), "restart": bool(restarted.search("pytest")), "invalidated": fact["id"] in invalidated}
    elif identifier == "case_05_office_files":
        seed = yaml.safe_load((workspace / "content.yaml").read_text(encoding="utf-8"))
        runtime.load_tools(["docx_create", "docx_read", "docx_edit", "xlsx_create", "xlsx_read", "xlsx_edit"])
        original_doc = require_ok(call(runtime, "docx_create", {"path": "original.docx", "title": seed["title"], "paragraphs": seed["paragraphs"]}))
        original_xlsx = require_ok(call(runtime, "xlsx_create", {"path": "original.xlsx", "sheets": seed["sheets"]}))
        require_ok(call(runtime, "docx_read", {"path": original_doc["path"]}))
        edited_doc = require_ok(call(runtime, "docx_edit", {"source": original_doc["path"], "path": "edited.docx", "find": "Replace", "replace": "Updated"}))
        edited_xlsx = require_ok(call(runtime, "xlsx_edit", {"source": original_xlsx["path"], "path": "edited.xlsx", "cells": {"Data!B2": "=1+1"}}))
        runtime.catalog.reset(); runtime.load_tools(["pptx_create", "pptx_read", "pdf_create", "pdf_read"])
        deck = require_ok(call(runtime, "pptx_create", {"path": "created.pptx", "slides": [{"title": seed["title"], "bullets": ["Verified"]}]}))
        pdf = require_ok(call(runtime, "pdf_create", {"path": "created.pdf", "lines": [seed["title"], "Verified"]}))
        runtime.load_tools(["docx_read"])
        checks = {"docx": require_ok(call(runtime, "docx_read", {"path": edited_doc["path"]}))["paragraphs"][-1]["text"].startswith("Updated"), "xlsx": edited_xlsx["verified"], "pptx": bool(require_ok(call(runtime, "pptx_read", {"path": deck["path"]}))["slides"]), "pdf": require_ok(call(runtime, "pdf_read", {"path": pdf["path"]}))["page_count"] == 1, "originals_preserved": Path(original_doc["path"]).is_file() and Path(original_xlsx["path"]).is_file()}
    else:
        raise ValueError("未知案例")
    return {"status": "passed" if all(checks.values()) else "failed", "assertions": checks, **runtime_report(store, runtime)}


def run_deepseek_case(case, workspace, run_root):
    from genesisai.app.cli import load_environment
    load_environment(workspace=workspace)
    if case["id"] == "case_01_code_fix":
        subprocess.run(["git", "init"], cwd=workspace, check=True, capture_output=True)
        subprocess.run(["git", "add", "."], cwd=workspace, check=True, capture_output=True)
        subprocess.run(["git", "-c", "user.name=GenesisAI", "-c", "user.email=simulation@example.invalid", "commit", "-m", "fixture"], cwd=workspace, check=True, capture_output=True)
    elif case["id"] == "case_05_office_files":
        from docx import Document
        from openpyxl import Workbook
        seed = yaml.safe_load((workspace / "content.yaml").read_text(encoding="utf-8"))
        document = Document()
        document.add_heading(seed["title"], level=1)
        for paragraph in seed["paragraphs"]:
            document.add_paragraph(paragraph)
        document.save(workspace / "source.docx")
        workbook = Workbook()
        for index, (name, rows) in enumerate(seed["sheets"].items()):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = name
            for row in rows:
                sheet.append(row)
        workbook.save(workspace / "source.xlsx")
    store, runtime, _ = make_runtime(workspace, run_root)
    model, remote = build_model(None)
    if not remote: raise RuntimeError("P7 DeepSeek 模式要求远程 OpenAI 兼容模型")
    original_tests = {
        path.relative_to(workspace).as_posix(): path.read_bytes()
        for path in workspace.rglob("test_*.py")
    }
    result = Runner(model, runtime, max_rounds=24, seconds=600).start(case["prompt"])
    if case["id"] == "case_03_web_research" and result["status"] == "completed":
        follow = Runner(model, runtime, max_rounds=12, seconds=300).start("继续详细说明刚才提到的一个产品及其主要优势")
        result = {"status": follow["status"], "answer": follow.get("answer"), "first_status": result["status"]}
    checks = {"runner_completed": result["status"] == "completed"}
    if case["id"] == "case_01_code_fix":
        tested = subprocess.run([sys.executable, "-m", "pytest", "-q"], cwd=workspace, capture_output=True, timeout=60)
        checks.update(
            tests_pass=tested.returncode == 0,
            tests_unchanged=all((workspace / name).read_bytes() == payload for name, payload in original_tests.items()),
            bug_fix_skill="bug_fix" in store.data["tool_runtime"].get("active_skills", []),
        )
    elif case["id"] == "case_02_file_management":
        windows = workspace / "windows.txt"
        payload = windows.read_bytes() if windows.is_file() else b""
        checks.update(
            bom_preserved=payload.startswith(b"\xef\xbb\xbf"),
            crlf_preserved=b"\r\n" in payload,
            file_organized=(workspace / "organized" / "notes.md").is_file(),
        )
    elif case["id"] == "case_03_web_research":
        checks.update(sources_fetched=bool(store.data.get("sources")), follow_up_completed=result["status"] == "completed")
    elif case["id"] == "case_04_memory_and_skill":
        reloaded_store = Store(run_root / "state", store.id, workspace=workspace)
        memory = MemoryManager(workspace)
        remembered = memory.search("pytest", type="command")
        sourced = [item for item in memory.search(type="project") if item.get("source_path")]
        invalidated = []
        if sourced:
            project_file = workspace / "project.toml"
            project_file.write_text(project_file.read_text(encoding="utf-8") + "# source changed\n", encoding="utf-8")
            invalidated = memory.validate_sources()
        checks.update(
            command_remembered=bool(remembered),
            skill_used="project_check" in reloaded_store.data["tool_runtime"].get("active_skills", []),
            restart_memory=bool(MemoryManager(workspace).search("pytest", type="command")),
            sourced_fact_invalidated=bool(sourced) and sourced[0]["id"] in invalidated,
        )
    elif case["id"] == "case_05_office_files":
        suffixes = {Path(item["path"]).suffix.casefold() for item in store.data.get("artifacts", {}).values()}
        called = {item.get("call", {}).get("name") for item in store.data.get("calls", {}).values()}
        checks.update(
            formats_created={".docx", ".xlsx", ".pptx", ".pdf"}.issubset(suffixes),
            limited_edits={"docx_edit", "xlsx_edit"}.issubset(called),
            formats_reopened={"docx_read", "xlsx_read", "pptx_read", "pdf_read"}.issubset(called),
            sources_preserved=(workspace / "source.docx").is_file() and (workspace / "source.xlsx").is_file(),
        )
    return {"status": "passed" if all(checks.values()) else "failed", "runner_status": result["status"], "assertions": checks, "answer": (result.get("answer") or "")[:4000], "model": {"provider": type(model).__name__, "name": getattr(model, "model", None)}, **runtime_report(store, runtime)}


def run_all(dataset, output, *, mode="offline", selected=None):
    dataset, output = Path(dataset).resolve(), Path(output).resolve(); output.mkdir(parents=True, exist_ok=True)
    started = time.time(); results = []
    for case, fixture in load_dataset(dataset):
        if selected and case["id"] not in selected: continue
        case_started = time.time()
        before = inventory(fixture)
        case_root = output / case["id"]; workspace = case_root / "workspace"
        if case_root.exists():
            if not case_root.resolve().is_relative_to(output):
                raise ValueError("仿真清理目标越界")
            def remove_readonly(function, path, _error):
                os.chmod(path, stat.S_IWRITE)
                function(path)
            shutil.rmtree(case_root, onexc=remove_readonly)
        shutil.copytree(fixture, workspace)
        try:
            detail = run_offline_case(case, workspace, case_root) if mode == "offline" else run_deepseek_case(case, workspace, case_root)
        except Exception as exc:
            detail = {"status": "failed", "error": f"{type(exc).__name__}: {exc}"}
        detail.update(
            case_id=case["id"], mode=mode,
            started_at=case_started, finished_at=time.time(),
            declared_assertions=case["assertions"], forbidden_rules=case["forbidden"],
            initial_inventory=before,
            fixture_unchanged=before == inventory(fixture), final_inventory=inventory(workspace),
        )
        if not detail["fixture_unchanged"]: detail["status"] = "failed"
        results.append(detail)
    report = {"mode": mode, "platform": platform.platform(), "python": platform.python_version(), "started_at": started, "finished_at": time.time(), "status": "passed" if results and all(item["status"] == "passed" for item in results) else "failed", "results": results}
    atomic_json(output / f"report-{mode}.json", report)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description="GenesisAI P7 产品仿真")
    parser.add_argument("--dataset", type=Path, default=Path(__file__).resolve().parents[3] / "tests" / "product_acceptance_dataset")
    parser.add_argument("--output", type=Path, default=Path.cwd() / ".genesis" / "simulation")
    parser.add_argument("--mode", choices=["offline", "deepseek"], default="offline")
    parser.add_argument("--case", action="append", dest="cases")
    args = parser.parse_args(argv)
    report = run_all(args.dataset, args.output, mode=args.mode, selected=set(args.cases or []))
    print(json.dumps({"status": report["status"], "mode": report["mode"], "cases": [{"id": x["case_id"], "status": x["status"]} for x in report["results"]]}, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())



