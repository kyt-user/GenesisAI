"""参考 xiaoerAI 文件格式模块实现的有界文档提取。"""
import csv
import io
import json
import os
import zipfile
from pathlib import Path

from genesisai.shared.security import ToolError
from genesisai.shared.filetypes import TEXT_SUFFIXES

MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT = 200_000
SUFFIXES = {*TEXT_SUFFIXES, '.docx', '.xlsx', '.pdf'}


def read_text_document(path):
    """读取普通文本并返回内容、实际编码、主换行和 BOM 状态。"""
    path = Path(path)
    raw = path.read_bytes()
    if len(raw) > MAX_BYTES:
        raise ToolError('size_limit', '文件超过 8 MiB 上限')
    if b'\x00' in raw[:8192]:
        raise ToolError('unsupported_format', '目标是二进制文件，不能作为文本修改')
    bom = raw.startswith(b'\xef\xbb\xbf')
    encoding = 'utf-8-sig' if bom else 'utf-8'
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = 'gb18030'
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            raise ToolError('unsupported_encoding', '文件编码不受支持') from exc
    crlf = text.count('\r\n')
    lf = text.count('\n') - crlf
    newline = '\r\n' if crlf > lf else '\n'
    return text, encoding, newline, bom


def encode_text_document(text, encoding, bom):
    target_encoding = 'utf-8' if encoding == 'utf-8-sig' else encoding
    payload = text.encode(target_encoding)
    if bom and target_encoding == 'utf-8':
        payload = b'\xef\xbb\xbf' + payload
    return payload


def extract(path):
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ToolError('size_limit', '文件超过 8 MiB 上限')
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        try:
            text, _, _, _ = read_text_document(path)
        except ToolError as exc:
            raise ToolError('unsupported_format', '不支持的格式') from exc
        return text[:MAX_TEXT], len(text) > MAX_TEXT
    truncated = False
    if suffix in {'.docx', '.xlsx'}:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 32 * 1024 * 1024:
                raise ToolError('size_limit', '文档解压后过大')
    if suffix in TEXT_SUFFIXES:
        text, _, _, _ = read_text_document(path)
        if suffix == '.json':
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        elif suffix == '.csv':
            text = '\n'.join('\t'.join(row) for row in csv.reader(io.StringIO(text)))
    elif suffix == '.docx':
        from docx import Document
        document = Document(path)
        parts = [p.text for p in document.paragraphs[:3000]]
        truncated = len(document.paragraphs) > 3000 or len(document.tables) > 50
        for table in document.tables[:50]:
            truncated |= len(table.rows) > 200
            parts.extend('\t'.join(c.text for c in row.cells[:50]) for row in table.rows[:200])
        text = '\n'.join(parts)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts = []
        try:
            truncated = len(workbook.sheetnames) > 10
            for sheet in list(workbook)[:10]:
                parts.append('## ' + sheet.title)
                truncated |= (sheet.max_row or 0) > 1000 or (sheet.max_column or 0) > 50
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 1000, 1000), max_col=min(sheet.max_column or 50, 50), values_only=True):
                    parts.append('\t'.join('' if v is None else str(v) for v in row))
        finally:
            workbook.close()
        text = '\n'.join(parts)
    else:
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ToolError('encrypted', 'PDF 已加密')
        truncated = len(reader.pages) > 50
        text = '\n'.join(f'[page {i + 1}]\n{page.extract_text() or ""}' for i, page in enumerate(reader.pages[:50]))
        if not any((p.extract_text() or '').strip() for p in reader.pages[:50]):
            raise ToolError('no_text', '没有可读正文；可能是扫描 PDF，第一版不支持 OCR')
    if not text.strip():
        raise ToolError('no_text', '资料没有可读正文')
    return text[:MAX_TEXT], bool(truncated or len(text) > MAX_TEXT)


def walk(access, path, cap=1000):
    root = access.read(path)
    count = 0
    for directory, folders, files in os.walk(root, followlinks=False):
        folders[:] = [f for f in sorted(folders) if f not in {'.git', '.venv', 'node_modules', '__pycache__'} and not Path(directory, f).is_symlink() and not getattr(Path(directory, f), 'is_junction', lambda: False)()]
        for name in sorted(files):
            count += 1
            if count > cap:
                return
            candidate = Path(directory, name)
            try:
                yield access.read(str(candidate))
            except (ToolError, OSError):
                continue

